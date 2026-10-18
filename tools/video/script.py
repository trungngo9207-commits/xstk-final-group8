"""B1 & B2 step-by-step guide video: content and choreography."""
import os, sys, functools
import openpyxl
import pandas as pd
from PIL import Image, ImageDraw

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from engine import *  # noqa

# ------------------------------------------------------------------ source data
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
WORKBOOK = os.path.join(REPO, "workbooks", "Group_8_local_snapshot_2026-10-11.xlsx")
wbf = openpyxl.load_workbook(WORKBOOK)["Part B"]
wbv = openpyxl.load_workbook(WORKBOOK, data_only=True)["Part B"]
SRC = {}
for row in wbf.iter_rows(min_row=1, max_row=60, max_col=12):
    for c in row:
        if c.value is None:
            continue
        f = c.value
        if isinstance(f, openpyxl.worksheet.formula.ArrayFormula):
            f = f.text.replace("TRUE()", "TRUE")
        SRC[c.coordinate] = (f, fmt(wbv[c.coordinate].value))

BOLD = {"A1", "A3", "A19", "A25", "A34", "A41", "K41", "B20", "C20", "A52", "B52", "C52", "D52", "E52", "F52", "G52", "H52"}


def cellof(ref):
    f, v = SRC[ref]
    if isinstance(f, str) and f.startswith("="):
        return Cell(v, f, bold=ref in BOLD)
    return Cell(v if v != "" else str(f), None, bold=ref in BOLD)


def T(ref, **kw):
    """Type the workbook's own content of `ref`."""
    f, v = SRC[ref]
    return TypeCell(ref, f if isinstance(f, str) else str(f), v, bold=ref in BOLD, **kw)


def P(*refs, gap=0.12, select=False):
    return Pop([(r, cellof(r)) for r in refs], gap=gap, select=select)


def col(rng):
    return cells_in(rng)


wine = pd.read_excel(os.path.join(REPO, "data", "Wine_Dataset.xlsx"))

partB = Sheet("Part B", {"A": 380, "B": 170, "C": 140, "D": 110, "E": 140, "F": 110, "G": 110, "H": 100,
                         "I": 90, "J": 90, "K": 165, "L": 165}, 100)
partB.cells["A1"] = cellof("A1")
data = Sheet("Data", {"A": 60, "B": 80, "C": 90, "D": 75, "E": 65, "F": 120, "U": 170, "V": 200}, 100,
             hidden=[n2col(i) for i in range(7, 21)])
for j, name in enumerate(["ID", "Price", "Cases", "Score", "Age", "Napa valley"]):
    data.cells[f"{n2col(j + 1)}1"] = Cell(name, bold=True)
for i, rr in wine.iterrows():
    r = i + 2
    for j, name in enumerate(["ID", "Price", "Cases", "Score", "Age", "Napa valley"]):
        data.cells[f"{n2col(j + 1)}{r}"] = Cell(fmt(rr[name]))


def uv_cells(r):
    napa_ = int(wine.loc[r - 2, "Napa valley"])
    price_ = fmt(wine.loc[r - 2, "Price"])
    return [(f"U{r}", Cell(price_ if napa_ == 1 else "", f'=IF(F{r}=1,B{r},"")', num=True)),
            (f"V{r}", Cell(price_ if napa_ == 0 else "", f'=IF(F{r}=0,B{r},"")', num=True))]


wb = Workbook([Sheet("Part A", {}), partB, Sheet("Part C", {}), data], "Data", "Group_8.xlsx")
V = Video(wb)

# ------------------------------------------------------------------ slide helpers
CW, CH = W, 870


def base(eyebrow, title):
    img = Image.new("RGB", (CW, CH), BG)
    d = ImageDraw.Draw(img, "RGBA")
    if eyebrow:
        draw_text(d, 96, 78, eyebrow.upper(), "ui-b", 20, BLUE)
    draw_text(d, 96, 140, title, "ui-b", 48, INK)
    return img, d


def para(d, x, y, text, size=28, color=INK2, maxw=800, key="ui", lh=1.38):
    for ln in wrap(text, key, size, maxw):
        draw_text(d, x, y + size, ln, key, size, color)
        y += size * lh
    return y


def rich(d, x, y, text, size=26, maxw=800, lh=1.4, color=INK2):
    """Word-wrap text where **bold** spans are drawn bold in primary ink."""
    words, bold, prev_space = [], False, True
    for part in text.split("**"):
        toks = part.split(" ")
        for i, w_ in enumerate(toks):
            if w_:
                glue = i == 0 and not prev_space  # no space between e.g. **bold** and "."
                words.append((w_, bold, glue))
        prev_space = part.endswith(" ") or part == ""
        bold = not bold
    lines, cur, cw = [], [], 0
    sp = text_w(" ", "ui", size)
    for w_, b, glue in words:
        ww = text_w(w_, "ui-b" if b else "ui", size)
        if cur and cw + sp + ww > maxw and not glue:
            lines.append(cur)
            cur, cw = [], 0
        cw += (sp if cur and not glue else 0) + ww
        cur.append((w_, b, ww, glue))
    if cur:
        lines.append(cur)
    for ln in lines:
        xx = x
        for i, (w_, b, ww, glue) in enumerate(ln):
            if glue and i > 0:
                xx -= sp
            draw_text(d, xx, y + size, w_, "ui-b" if b else "ui", size, INK if b else color)
            xx += ww + sp
        y += size * lh
    return y


def card(d, box, fill=CARD, outline=AXIS, accent=None):
    x0, y0, x1, y1 = box
    rrect(d, (x0 + 2, y0 + 4, x1 + 2, y1 + 4), 16, fill=(0, 0, 0, 18))
    rrect(d, box, 16, fill=fill, outline=outline, width=1)
    if accent:
        d.rounded_rectangle((x0, y0, x0 + 8, y1), radius=4, fill=accent)


def paste_chart(img, path, x, y, w):
    im = load_img(path)
    sc = w / im.width
    im2 = im.resize((int(im.width * sc), int(im.height * sc)), Image.LANCZOS)
    img.paste(im2, (x, y))
    return im2.height


def S(fn):
    return functools.lru_cache(maxsize=None)(fn)


@S
def s_title():
    img = Image.new("RGB", (CW, CH), BG)
    d = ImageDraw.Draw(img, "RGBA")
    draw_text(d, 140, 250, "FINAL PROJECT · WINE PRICES AND REGIONAL PRICE PREMIA", "ui-b", 24, BLUE)
    draw_text(d, 136, 360, "Part B1 & B2, step by step", "ui-b", 84, INK)
    para(d, 140, 400, "From one average to a regression on one dummy: built in Excel, explained from zero.",
         34, INK2, 1500)
    d.line((140, 540, 600, 540), fill=AXIS, width=2)
    draw_text(d, 140, 600, "Group 8  ·  TOAE203 Statistics and Probability in Economics and Business", "ui", 26, MUTED)
    draw_text(d, 140, 645, "Pause any time: every formula is shown in the formula bar.", "ui", 26, MUTED)
    return img


LADDER = [
    ("B1", "No X", "Price = β₀ + ε", "mean · confidence interval · one-sample t"),
    ("B2", "One dummy X", "Price = β₀ + β₁·Napa + ε", "two-sample t (Welch, pooled) · rank test"),
    ("B3", "One X, 6 groups", "ln Price = β₀ + region dummies + ε", "one-way ANOVA"),
    ("B4", "One continuous X", "ln Price = β₀ + β₁·Score + ε", "correlation · simple regression"),
    ("B5", "Several X", "Price = β₀ + β₁X₁ + … + ε", "multiple regression"),
    ("B6", "Categorical Y", "Region × Variety", "crosstab · χ² test"),
]


@S
def s_ladder():
    img, d = base("The big picture", "B1 and B2 are the first two rungs of the model ladder")
    y = 210
    for i, (code, xs, eq, names) in enumerate(LADDER):
        on = i < 2
        card(d, (96, y, 1824, y + 88), fill=CARD if on else (236, 235, 231), outline=BLUE if on else GRID,
             accent=BLUE if on else None)
        colr = INK if on else MUTED
        draw_text(d, 130, y + 44, code, "ui-b", 32, BLUE if on else MUTED, anchor="lm")
        draw_text(d, 230, y + 44, xs, "ui-b", 26, colr, anchor="lm")
        draw_text(d, 560, y + 44, eq, "ui", 28, colr, anchor="lm")
        draw_text(d, 1790, y + 44, names, "ui", 23, INK2 if on else MUTED, anchor="rm")
        y += 104
    draw_text(d, 96, 850, "Your part: the two highlighted rungs. Each later rung adds one thing to the same equation.",
              "ui", 24, INK2)
    return img


@S
def s_tasks(k):
    img, d = base("What the project asks", "B1 (5 points) and B2 (8 points)")
    boxes = [(96, 210, 880, 820), (944, 210, 1824, 820)]
    content = [
        ("B1 · Continuous Y, no X", "Price = β₀ + ε",
         ["Compute the **mean of Price** and its **95% confidence interval**.",
          "Test **H₀: μ = 25** against **H₁: μ ≠ 25**.",
          "Explain in one sentence why this mean is the estimate of **β₀**."]),
        ("B2 · Continuous Y, one dummy X", "Price = β₀ + β₁·Napa + ε",
         ["**1.** Compare mean Price of Napa vs. all other wines with **Welch's t-test** (unequal variances).",
          "**2.** Regress Price on the Napa dummy. Show **slope = difference of the two means**.",
          "**3.** Which two-sample t-test does the regression reproduce exactly, and why?",
          "**4.** Name a safer **rank-based test** (no need to run it)."]),
    ]
    for i in range(k):
        x0, y0, x1, y1 = boxes[i]
        title, eq, items = content[i]
        card(d, boxes[i], accent=BLUE)
        draw_text(d, x0 + 40, y0 + 62, title, "ui-b", 32, INK)
        draw_text(d, x0 + 40, y0 + 112, eq, "ui", 28, BLUE)
        y = y0 + 150
        for it in items:
            y = rich(d, x0 + 40, y, it, 26, (x1 - x0) - 80) + 16
    return img


@S
def s_rules(k):
    img, d = base("Rule for every hypothesis test", "Four steps, every time")
    steps = [("1", "Hypotheses", "State H₀ and H₁."),
             ("2", "Test statistic", "Name it, and the Excel formula or output it comes from."),
             ("3", "p-value", "Give the p-value (or critical value) at α = 0.05."),
             ("4", "Decision", "Reject or not, and what it means for wine prices.")]
    x = 96
    for i, (n, t, s) in enumerate(steps):
        on = i < k
        box = (x, 240, x + 410, 620)
        card(d, box, fill=CARD if on else (236, 235, 231), outline=BLUE if on else GRID)
        d.ellipse((x + 36, 276, x + 100, 340), fill=BLUE if on else AXIS)
        draw_text(d, x + 68, 308, n, "ui-b", 32, WHITE, anchor="mm")
        draw_text(d, x + 36, 410, t, "ui-b", 34, INK if on else MUTED)
        para(d, x + 36, 440, s, 26, INK2 if on else MUTED, 340)
        x += 442
    if k >= 5:
        card(d, (96, 680, 1824, 800), fill=(232, 240, 251), outline=BLUE)
        rich(d, 136, 712, "Every number must be **computed in the workbook with a visible formula**. "
             "Raw output pasted without interpretation earns no credit.", 28, 1640)
    return img


@S
def s_beta0(k):
    img, d = base("B1 · last question", "Why is the mean the estimate of β₀?")
    y = 230
    blocks = [
        "The model **Price = β₀ + ε** has no X. It predicts the **same number β₀ for every wine**; "
        "ε is each wine's error: its price minus that prediction.",
        "Least squares (OLS) picks the β₀ that makes the **sum of squared errors Σ(Price − β₀)² as small as possible**.",
        "Setting the derivative to zero: −2·Σ(Price − β₀) = 0, so **β₀ = ΣPrice / n = x̄**. "
        "The chart shows the same thing: the U-shaped curve bottoms out at the mean, **$29.58**.",
    ]
    for i, b in enumerate(blocks[: min(k, 3)]):
        y = rich(d, 96, y, b, 28, 820) + 26
    if k >= 3:
        card(d, (980, 200, 1824, 760))
        paste_chart(img, "charts/sse.png", 1000, 215, 804)
    return img


@S
def s_b2intro(k):
    img, d = base("B2 · one dummy X", "Napa Valley wines vs. all the others")
    card(d, (96, 200, 1000, 840))
    paste_chart(img, "charts/strip_full.png", 112, 214, 872)
    y = 230
    items = ["The Napa dummy splits the 806 wines into **two groups: 287 Napa wines and 519 other wines.**",
             "Each dot is one wine. Napa wines cost **more on average** ($43.06 vs. $22.13)…",
             "…and their prices are **far more spread out**: a few cost over $300. Remember this, it matters for B2.3."]
    for it in items[:k]:
        y = rich(d, 1060, y, it, 29, 760) + 28
    return img


@S
def s_regression(k):
    img, d = base("B2.2 · the same comparison as a regression", "Price = β₀ + β₁ · Napa + ε")
    y = 230
    if k >= 1:
        card(d, (96, y, 900, y + 250))
        draw_text(d, 130, y + 60, "Not Napa  (Napa = 0)", "ui-b", 28, ORANGE)
        draw_text(d, 130, y + 105, "predicted price = β₀", "ui", 30, INK)
        draw_text(d, 130, y + 170, "Napa  (Napa = 1)", "ui-b", 28, BLUE)
        draw_text(d, 130, y + 215, "predicted price = β₀ + β₁", "ui", 30, INK)
    if k >= 2:
        rich(d, 96, 530, "OLS makes each prediction the **mean of its own group** (just like B1). So **β₀ = mean of the "
             "other wines** and **β₁ = mean(Napa) − mean(other)**: the difference between the two means.", 28, 800)
    if k >= 3:
        card(d, (960, 190, 1824, 840))
        paste_chart(img, "charts/strip_reg.png", 976, 204, 832)
    return img


@S
def s_pooled(k):
    img, d = base("B2.3 · same difference, different standard error", "Which t-test does the regression reproduce?")
    cols = [("Welch t-test", "T.TEST type 3", ["20.93", "3.115", "6.72", "306.3", "8.9E-11"], ORANGE),
            ("Regression slope", "LINEST / ToolPak", ["20.93", "2.403", "8.71", "804", "1.7E-17"], BLUE),
            ("Pooled t-test", "T.TEST type 2", ["20.93", "2.403", "8.71", "804", "1.7E-17"], BLUE)]
    rows = ["Difference in means ($)", "Standard error", "t statistic", "Degrees of freedom", "p-value (two-sided)"]
    x0 = 96
    for i, r in enumerate(rows):
        yy = 330 + i * 78
        draw_text(d, x0, yy, r, "ui", 28, INK2, anchor="lm")
        d.line((x0, yy + 39, 1824, yy + 39), fill=GRID, width=1)
    for j, (name, sub, vals, colr) in enumerate(cols[:k]):
        cx = 820 + j * 340
        draw_text(d, cx + 150, 240, name, "ui-b", 30, INK, anchor="mm")
        draw_text(d, cx + 150, 280, sub, "ui", 22, MUTED, anchor="mm")
        d.rounded_rectangle((cx + 40, 296, cx + 260, 300), radius=2, fill=colr)
        for i, v in enumerate(vals):
            draw_text(d, cx + 150, 330 + i * 78, v, "ui-b", 32, INK, anchor="mm")
    if k >= 3:
        d.rounded_rectangle((1150, 214, 1830, 720), radius=16, outline=BLUE, width=3)
        rich(d, 96, 760, "The regression slope test **is** the pooled (equal-variance) two-sample t-test: "
             "same standard error, same t, same df, same p-value.", 28, 1700)
    return img


@S
def s_rank(k):
    img, d = base("B2.4 · a safer check", "Price is skewed, so use a rank-based test")
    card(d, (96, 200, 1000, 760))
    paste_chart(img, "charts/hist.png", 112, 214, 872)
    y = 220
    if k >= 1:
        rich(d, 1060, y, "**Wilcoxon rank-sum test**, also called the **Mann–Whitney U test**: the rank version of "
             "the two-sample t-test.", 30, 760)
    if k >= 2:
        prices = ["$12", "$18", "$25", "$60", "$440"]
        yy = 400
        draw_text(d, 1060, yy, "price", "ui", 24, MUTED, anchor="lm")
        draw_text(d, 1060, yy + 70, "rank", "ui", 24, MUTED, anchor="lm")
        for i, p in enumerate(prices):
            xx = 1180 + i * 125
            rrect(d, (xx, yy - 26, xx + 105, yy + 26), 10, fill=WHITE, outline=AXIS)
            draw_text(d, xx + 52, yy, p, "ui-b", 26, INK, anchor="mm")
            d.line((xx + 52, yy + 30, xx + 52, yy + 42), fill=AXIS, width=2)
            rrect(d, (xx + 18, yy + 46, xx + 86, yy + 94), 10, fill=(232, 240, 251), outline=BLUE)
            draw_text(d, xx + 52, yy + 70, str(i + 1), "ui-b", 28, INK, anchor="mm")
        rich(d, 1060, 520, "Each price is replaced by its **rank**. The $440 bottle is just **rank 5**: "
             "an extreme price cannot drag the result.", 27, 760)
    if k >= 3:
        rich(d, 1060, 650, "You only need to **name** it and say why it is safer. No need to run it.", 27, 760)
    return img


@S
def s_report_b1():
    img, d = base("Writing it up · B1", "What goes in the report for B1")
    card(d, (96, 200, 1824, 830), accent=BLUE)
    y = 235
    items = [
        "**Key numbers (Part B!B4:B15):** n = 806, mean Price = **$29.58**, s = $34.15, SE = $1.20, "
        "t(0.025, 805) = 1.963, **95% CI = [$27.22 ; $31.94]**.",
        "**Test:** H₀: μ = 25 vs H₁: μ ≠ 25. One-sample t = (29.58 − 25) / 1.20 = **3.81** (B12). "
        "p = T.DIST.2T(3.81, 805) = **0.00015** (B13) < 0.05, so **reject H₀**.",
        "**Meaning:** the average price of these California wines is significantly above $25. "
        "Consistent with the CI: $25 lies outside [$27.22 ; $31.94].",
        "**Why β₀ = x̄:** in Price = β₀ + ε, OLS chooses β₀ to minimise Σ(Price − β₀)²; that sum is smallest "
        "at the sample mean, so the estimate of β₀ is x̄ = $29.58.",
    ]
    for it in items:
        y = rich(d, 136, y, it, 28, 1640) + 22
    return img


@S
def s_report_b2():
    img, d = base("Writing it up · B2", "What goes in the report for B2")
    card(d, (96, 200, 1824, 845), accent=BLUE)
    y = 225
    items = [
        "**B2.1** H₀: μ(Napa) = μ(other) vs H₁: μ(Napa) ≠ μ(other). Napa: n = 287, mean **$43.06**; other: n = 519, "
        "mean **$22.13**. Welch t = **6.72** (B28), df ≈ 306, p = **8.9×10⁻¹¹** (B30, T.TEST type 3) < 0.05 → "
        "**reject H₀**: Napa wines cost about **$20.93** more per bottle on average.",
        "**B2.2** Fitted line: Price = **22.13 + 20.93·Napa** (LINEST, K42:L46). Intercept = mean of other wines; "
        "slope = 43.06 − 22.13 = 20.93 = difference of means (check cells B56, B57 = 0).",
        "**B2.3** Slope t = **8.71** = **pooled (equal-variance) t-test**, df = 804 (B37, B58 = 0). OLS assumes one "
        "common error variance for all wines, so its SE is the pooled SE (2.40). Welch keeps two variances (SE 3.11).",
        "**B2.4** Price is strongly right-skewed (A5), so a safer check is the **Wilcoxon rank-sum / Mann–Whitney U "
        "test**, which uses ranks and is not driven by a few very expensive bottles.",
    ]
    for it in items:
        y = rich(d, 136, y, it, 26, 1640, lh=1.36) + 18
    return img


QA = [
    ("Why STDEV.S and not STDEV.P?",
     "Our 806 wines are a sample. STDEV.S divides by n − 1, which corrects for estimating the mean from the same data."),
    ("Why 1.963 and not 1.96?",
     "σ is unknown, so we use s and the t distribution with n − 1 = 805 df. With n this large, t is almost z."),
    ("What does p = 0.00015 mean?",
     "If μ really were 25, a t at least this extreme would occur in about 0.015% of samples. It is not P(H₀ is true)."),
    ("Does B2 prove the Napa name causes higher prices?",
     "No. It is a raw comparison. Napa wines may also score higher, be more often Cabernet, or age longer. "
     "B5 holds those constant."),
    ("Price is not normal. Are the t-tests still valid?",
     "Largely yes: with n = 806 the central limit theorem makes the sample means nearly normal. "
     "Mann–Whitney is a robustness check."),
    ("Why does the regression give the pooled t, not Welch's?",
     "OLS uses one error variance for every wine: SE of regression = 32.67, and 32.67² = 1,067 = the pooled variance."),
]


QA_SAY = [
    ("Question: why STDEV dot S, and not STDEV dot P?",
     "Because our 806 wines are a sample. STDEV dot S divides by n minus 1, which corrects for estimating the mean "
     "from the same data."),
    ("Why 1.963, and not 1.96?",
     "Because sigma, the true standard deviation, is unknown. We use s instead, so we need the t distribution with "
     "805 degrees of freedom. With a sample this large, t is almost the same as z."),
    ("What does a p-value of 0.00015 mean?",
     "If mu really were 25, a t statistic at least this extreme would happen in only about 0.015 percent of samples. "
     "It is not the probability that H zero is true."),
    ("Does B2 prove that the Napa name causes higher prices?",
     "No. It is a raw comparison. Napa wines may also score higher, be more often Cabernet, or be aged longer. "
     "The multiple regression in B5 holds those things constant."),
    ("Price is not normal. Are the t-tests still valid?",
     "Largely, yes. With 806 wines, the central limit theorem makes the sample means nearly normal. "
     "The Mann-Whitney test is a robustness check."),
    ("Why does the regression give the pooled t, and not Welch's t?",
     "Because least squares uses one error variance for every wine. The standard error of the regression is 32.67, "
     "and 32.67 squared is 1,067, which is exactly the pooled variance."),
]


@S
def s_qa(page, k):
    img, d = base("Q&A practice", "Questions the lecturer may ask you" + (" (2/2)" if page else " (1/2)"))
    y = 215
    for i, (q, a) in enumerate(QA[page * 3: page * 3 + k]):
        card(d, (96, y, 1824, y + 190), accent=BLUE)
        draw_text(d, 136, y + 52, "Q  " + q, "ui-b", 30, INK)
        para(d, 136, y + 78, a, 26, INK2, 1640)
        y += 214
    return img


@S
def s_end():
    img, d = base("Recap", "You can now do B1 and B2")
    items = ["n, mean, s, SE with COUNT, AVERAGE, STDEV.S, s/√n",
             "95% CI = x̄ ± T.INV.2T(0.05, n−1) × SE",
             "One-sample t, p-value with T.DIST.2T, decision with IF",
             "Why β₀ = x̄: least squares is minimised at the mean",
             "Welch t-test: own variances, T.TEST(…, 2, 3)",
             "LINEST on the dummy: intercept = mean(other), slope = difference",
             "Regression t = pooled t: one common error variance",
             "Skewed Price → Wilcoxon rank-sum / Mann–Whitney U"]
    y = 220
    for it in items:
        d.ellipse((100, y + 8, 128, y + 36), fill=BLUE)
        d.line((107, y + 22, 112, y + 28), fill=WHITE, width=3)
        d.line((112, y + 28, 121, y + 15), fill=WHITE, width=3)
        draw_text(d, 150, y + 32, it, "ui", 30, INK)
        y += 66
    return img


def slide(fn, *args):
    return lambda: fn(*args)


def note(*items):
    return list(items)


# ================================================================== SCRIPT
V.chapter("Welcome")
V.slide(slide(s_title), [
    "Hi! In this video we build parts B1 and B2 of the Group 8 final project together, one step at a time.",
    "You do not need any background. I explain every formula before we type it, and at the end you will know "
    "what to write in the report and how to answer the lecturer's questions.",
])

V.chapter("What B1 and B2 ask")
V.slide(slide(s_ladder), [
    "The project asks one big question: do buyers pay more for a bottle labelled Napa Valley?",
    "Part B climbs towards the answer like a ladder. Each step adds one thing to the same equation, and B1 and B2 "
    "are the first two steps.",
])
V.slide(slide(s_tasks, 1), [
    "In B1 we look at Price on its own, with no X. We compute the average price, put a 95% confidence interval "
    "around it, and test whether the true average could be $25.",
])
V.slide(slide(s_tasks, 2), [
    "In B2 we add one yes-or-no variable: is the wine from Napa Valley? We compare the two groups with a t-test, "
    "then show that a simple regression gives exactly the same comparison.",
])
for k, txt in [(1, "The project has one strict rule for every hypothesis test. One: state the null hypothesis, H₀, "
                   "and the alternative, H₁."),
               (2, "Two: name the test statistic, and the Excel formula it comes from."),
               (3, "Three: give the p-value, and compare it with alpha, 0.05."),
               (4, "Four: state the decision, and say what it means for wine prices."),
               (5, "And every number must be computed in the workbook with a visible formula. That is why we type "
                   "formulas instead of copying numbers.")]:
    V.slide(slide(s_rules, k), [txt])

# ------------------------------------------------------------------ data
V.chapter("Meet the data")
V.sheet([
    "Here is the Data sheet of our team workbook. Each row is one wine: 806 wines in total, so the data runs "
    "from row 2 down to row 807.",
    "To make everything fit on screen, I have hidden the columns we do not need today.",
], [Sync(1), badge("columns G to T hidden in this view", 585, 196, INK2)],
    note=note(("eyebrow", "The data"), ("title", "806 California wines, 2001 vintage"),
              ("text", "One row = one wine. Rows 2 to 807 hold the data; row 1 holds the column names."),
              ("kv", [("Rows of data", "806"), ("First data row", "2"), ("Last data row", "807")])))
V.sheet([
    "Column B is Price, the retail price of one bottle in US dollars. This is our Y: the variable we want to explain.",
], [clear_badges(), callout("B1:B23", BLUE, "Price  (Y)", "top")],
    note=note(("eyebrow", "Column B"), ("title", "Price = Y"),
              ("text", "Retail price per bottle, in US dollars. Everything in B1 and B2 is about this column.")))
V.sheet([
    "Column F is the Napa valley dummy. A dummy is a yes-or-no variable written as a number: 1 means the wine "
    "comes from Napa Valley, 0 means it does not.",
    "For B1 and B2 we only need these two columns: Price, and the Napa dummy.",
], [callout("F1:F23", ORANGE, "Napa dummy", "top")],
    note=note(("eyebrow", "Column F"), ("title", "Napa valley dummy"),
              ("kv", [("1", "wine is from Napa Valley"), ("0", "wine is from anywhere else")]),
              ("text", "A dummy turns a yes/no question into a number, so it can go into a formula or a regression.")))

# ------------------------------------------------------------------ B1 basics
V.chapter("B1 · Count, mean, spread")
V.sheet([
    "Now open the Part B sheet. First, type the labels in column A, so that every number has a name next to it.",
], [clear_callouts(), GotoSheet("Part B"), P("A3", "A4", "A5", "A6", "A7", "A8", "A9", "A10", "A11", "A12", "A13",
                                              "A14", "A15", gap=0.14)],
    note=note(("eyebrow", "B1 · setup"), ("title", "Labels first"),
              ("text", "Column A holds the names, column B the formulas. A reader (and the lecturer) can then trace "
                       "every number.")))
V.sheet([
    "In B4 we count the wines. Type equals COUNT, open a bracket, and give it the Price column of the Data sheet: "
    "Data, exclamation mark, B2 to B807.",
    "Press Enter. The answer is 806.",
    "The dollar signs lock the range. They have nothing to do with money: they just stop the range from moving if "
    "you copy the formula.",
], [T("B4")],
    note=note(("eyebrow", "B1 · step 1"), ("title", "n = number of wines"),
              ("code", "=COUNT(Data!$B$2:$B$807)"),
              ("text", "COUNT counts the cells that contain numbers."),
              ("text", "Data!  = look on the Data sheet.\n$B$2:$B$807 = the Price column, rows 2 to 807."),
              ("text", "$ locks the row/column so the range never shifts when copied.")))
V.sheet([
    "In B5, the mean. Type equals AVERAGE of the same range. The average bottle in our sample costs $29.58. "
    "We call it x̄.",
], [T("B5")],
    note=note(("eyebrow", "B1 · step 2"), ("title", "Mean x̄"), ("code", "=AVERAGE(Data!$B$2:$B$807)"),
              ("kv", [("x̄", "$29.58")]),
              ("text", "The sum of all prices divided by 806.")))
V.sheet([
    "In B6, the standard deviation, with STDEV.S. The S stands for sample: it divides by n minus 1, because our "
    "806 wines are a sample, not every wine in California.",
    "The answer is $34.15. That is large compared with the mean, so prices are very spread out.",
], [T("B6")],
    note=note(("eyebrow", "B1 · step 3"), ("title", "Standard deviation s"), ("code", "=STDEV.S(Data!$B$2:$B$807)"),
              ("kv", [("s", "$34.15")]),
              ("text", "How far a typical wine's price is from the mean. STDEV.S = sample version (divides by n − 1).")))
V.sheet([
    "In B7, the standard error: the standard deviation divided by the square root of n. Type equals B6, a slash, "
    "then SQRT of B4.",
    "The standard deviation says how much single wines differ. The standard error says how much the average "
    "itself would move if we drew another sample of 806 wines. Here it is only $1.20.",
], [T("B7")],
    note=note(("eyebrow", "B1 · step 4"), ("title", "Standard error SE = s / √n"), ("code", "=B6/SQRT(B4)"),
              ("kv", [("s (single wines)", "$34.15"), ("SE (the mean)", "$1.20")]),
              ("text", "The SE shrinks as n grows: a mean of 806 wines is much more stable than one wine.")))

# ------------------------------------------------------------------ CI
V.chapter("B1 · 95% confidence interval")
V.sheet([
    "A 95% confidence interval is the mean, plus or minus a critical value times the standard error. First we "
    "need the critical value.",
    "In B8, type equals T.INV.2T, open a bracket, 0.05, comma, B4 minus 1, close the bracket.",
    "0.05 is alpha: we leave 5% outside the interval, 2.5% in each tail. That is what 2T means: two tails. B4 minus 1 "
    "is the degrees of freedom, 805. The answer is 1.963.",
    "We use the t distribution because we only know the sample standard deviation, not the true one. With 806 "
    "wines, t is almost the normal value, 1.96.",
], [Sync(1), T("B8")],
    note=note(("eyebrow", "B1 · confidence interval"), ("title", "CI = x̄ ± t* × SE"), ("code", "=T.INV.2T(0.05,B4-1)"),
              ("kv", [("α (two tails)", "0.05"), ("df = n − 1", "805"), ("t*", "1.963")]),
              ("text", "T.INV.2T(α, df) gives the t value that leaves α/2 in each tail.")))
V.sheet([
    "In B9, the lower limit: B5 minus B8 times B7. In B10, the upper limit: B5 plus B8 times B7.",
    "So the 95% confidence interval for the mean price runs from $27.22 to $31.94.",
], [T("B9"), T("B10"), callout("B9:B10", BLUE, "95% CI", "right")],
    note=note(("eyebrow", "B1 · confidence interval"), ("title", "95% CI: $27.22 to $31.94"),
              ("code", "=B5-B8*B7\n=B5+B8*B7"), ("img", "charts/ci.png")))
V.sheet([
    "How to read it: if we repeated the sampling many times, 95% of the intervals built this way would contain "
    "the true mean price.",
    "It does not mean that 95% of wines cost between $27 and $32. In fact, fewer than 8% of our wines are in that range.",
], [],
    note=note(("eyebrow", "Read it correctly"), ("title", "What 95% means"),
              ("strong", "Correct: 95% of intervals built this way contain the true mean μ."),
              ("text", "Wrong: 95% of wines cost between $27.22 and $31.94. (Only 7.6% of our wines do.)")))
V.sheet([
    "Excel can check our work. In B15, CONFIDENCE.T of 0.05, B6 and B4 gives the half-width of the interval, 2.36. "
    "That is exactly the critical value times the standard error.",
], [T("B15"), clear_callouts()],
    note=note(("eyebrow", "B1 · check"), ("title", "Margin of error"), ("code", "=CONFIDENCE.T(0.05,B6,B4)"),
              ("kv", [("t* × SE = 1.963 × 1.203", "2.36"), ("CONFIDENCE.T", "2.36")]),
              ("text", "Same number, so the interval is right: 29.58 ± 2.36.")))

# ------------------------------------------------------------------ test
V.chapter("B1 · Test H₀: μ = 25")
V.sheet([
    "Now the hypothesis test. Step one, the hypotheses. H₀: the true mean price μ equals $25. H₁: μ is not equal to "
    "25. Type 25 in B11.",
], [T("B11")],
    note=note(("eyebrow", "Step 1 · hypotheses"), ("title", "H₀: μ = 25\nH₁: μ ≠ 25"),
              ("text", "Two-sided: we would reject if the mean were far above OR far below $25."),
              ("text", "μ is the true mean price of all such wines; x̄ = 29.58 is our estimate of it.")))
V.sheet([
    "Step two, the test statistic. In B12, type: open a bracket, B5 minus B11, close the bracket, divided by B7.",
    "This is the t statistic: signal divided by noise. The signal is how far our mean is from 25. The noise is the "
    "standard error. t equals 3.81: our mean sits 3.81 standard errors above $25.",
], [T("B12")],
    note=note(("eyebrow", "Step 2 · test statistic"), ("title", "One-sample t = (x̄ − μ₀) / SE"),
              ("code", "=(B5-B11)/B7"), ("kv", [("signal: x̄ − 25", "4.58"), ("noise: SE", "1.20"), ("t", "3.81")])))
V.sheet([
    "Step three, the p-value. In B13, type T.DIST.2T, with the absolute value of B12, and B4 minus 1 degrees of freedom.",
    "The p-value is 0.00015. It means: if the true mean really were $25, a t this far from zero would happen only "
    "about 15 times in 100,000 samples.",
], [T("B13")],
    note=note(("eyebrow", "Step 3 · p-value"), ("title", "p = 0.00015"), ("code", "=T.DIST.2T(ABS(B12),B4-1)"),
              ("img", "charts/tcurve.png"),
              ("text", "2T = both tails. ABS because T.DIST.2T needs a positive t.")))
V.sheet([
    "Step four, the decision. In B14, an IF formula: if B13 is less than 0.05, show Reject H0, otherwise show "
    "Do not reject H0. The cell says Reject H0.",
    "In wine terms: the average price of these California wines is significantly higher than $25.",
], [T("B14")],
    note=note(("eyebrow", "Step 4 · decision"), ("title", "p < 0.05 → reject H₀"),
              ("code", '=IF(B13<0.05,"Reject H0","Do not reject H0")'),
              ("strong", "The mean price is significantly above $25.")))
V.sheet([
    "Notice that this agrees with the confidence interval. $25 lies outside the interval from 27.22 to 31.94, "
    "and a two-sided test at 5% rejects exactly when the value lies outside the 95% interval.",
    "Finally, the team wrote the conclusion in one sentence in row 17. You will reuse it in the report.",
], [callout("B9:B10", BLUE, "95% CI", "right"), callout("B11", ORANGE, "H₀ value", "right"), Sync(1),
    clear_callouts(), P("A17")],
    note=note(("eyebrow", "Same answer, two ways"), ("title", "Test and interval agree"), ("img", "charts/ci.png"),
              ("text", "Two-sided test at α = 0.05 rejects ⇔ μ₀ is outside the 95% CI.")))

V.chapter("B1 · Why the mean is β₀")
V.slide(slide(s_beta0, 1), [
    "The last part of B1 asks why this mean is the estimate of β₀, in the model Price equals β₀ plus epsilon.",
    "This model has no X at all. It predicts the same number, β₀, for every wine. Epsilon is each wine's error: "
    "its price minus that prediction.",
])
V.slide(slide(s_beta0, 2), [
    "Regression, meaning ordinary least squares, picks β₀ to make the sum of squared errors as small as possible.",
])
V.slide(slide(s_beta0, 3), [
    "The chart shows that sum for every possible value of β₀. It is a U shape, and the bottom of the U is exactly "
    "at the mean, $29.58.",
    "So, in one sentence: with no X, least squares chooses the β₀ that minimises the squared errors, and that "
    "minimum is reached at the sample mean. Our estimate of β₀ is x̄, $29.58.",
])

# ------------------------------------------------------------------ B2 setup
V.chapter("B2 · Napa vs. the rest")
V.slide(slide(s_b2intro, 1), [
    "Now B2. The Napa dummy splits the wines into two groups: 287 Napa wines and 519 other wines.",
])
V.slide(slide(s_b2intro, 3), [
    "Each dot on this chart is one wine. Napa wines cost more on average.",
    "They are also much more spread out, with a few bottles above $300. Keep that second fact in mind: it matters later.",
])
V.sheet([
    "To compare the groups in Excel, we first make two helper columns on the Data sheet: one with only the Napa "
    "prices, and one with only the other prices.",
    "In U1, type the header: Price if Napa.",
], [GotoSheet("Data"), Sync(1), TypeCell("U1", "Price if Napa", "Price if Napa", bold=True)],
    note=note(("eyebrow", "B2 · helper columns"), ("title", "Split Price into two columns"),
              ("text", "Column U: price if the wine is from Napa, else empty.\nColumn V: price if it is not, else empty."),
              ("text", "COUNT, AVERAGE and VAR.S ignore empty text, so each column only describes its own group.")))
V.sheet([
    "In U2, type: equals IF, open a bracket, F2 equals 1, comma, B2, comma, two quotation marks, close the bracket.",
    "It reads: if this wine is from Napa, copy its price; otherwise leave the cell empty. Two quotation marks mean "
    "empty text.",
], [TypeCell("U2", '=IF(F2=1,B2,"")', "", num=True)],
    note=note(("eyebrow", "B2 · helper columns"), ("title", "Price if Napa"), ("code", '=IF(F2=1,B2,"")'),
              ("text", "IF(test, value if true, value if false)."),
              ("kv", [("F2 = 1 (Napa)", "copy B2"), ("otherwise", "\"\" = empty")])))
V.sheet([
    "In V1, type Price if not Napa. In V2, the opposite test: if F2 equals 0, copy the price, otherwise leave it empty.",
], [TypeCell("V1", "Price if not Napa", "Price if not Napa", bold=True),
    TypeCell("V2", '=IF(F2=0,B2,"")', "6", num=True)],
    note=note(("eyebrow", "B2 · helper columns"), ("title", "Price if not Napa"), ("code", '=IF(F2=0,B2,"")')))
V.sheet([
    "Now select U2 and V2 together, and double-click the small square at the bottom-right corner, the fill handle. "
    "Excel copies both formulas all the way down to row 807.",
], [Click("U2", sel="U2:V2"), FillDown("U2:V2", uv_cells, 807), badge("filled down to row 807", 950, 230, BLUE)],
    note=note(("eyebrow", "B2 · helper columns"), ("title", "Fill down"),
              ("text", "Double-click the fill handle (the small square at the corner of the selection). "
                       "The formulas adjust row by row: F3, B3, F4, B4…")))
V.sheet([
    "Scroll down and check: every price now appears in exactly one of the two columns. U for Napa wines, V for all "
    "the others.",
], [clear_badges(), Do(lambda w: (setattr(w, "sel", None), setattr(w, "cur", None))), Scroll(top=196),
    callout("F201:F217", ORANGE, None), callout("U201:V217", BLUE, None)],
    note=note(("eyebrow", "Check"), ("title", "Each wine lands in one column"),
              ("kv", [("Napa = 1", "price in U"), ("Napa = 0", "price in V")])))

# ------------------------------------------------------------------ Welch
V.chapter("B2.1 · Welch t-test")
V.sheet([
    "Back on Part B, under B1, type the labels for B2: n, mean price and variance, with one column for Napa and "
    "one for the other wines.",
], [clear_callouts(), GotoSheet("Part B"), Scroll(top=17), P("A19", "B20", "C20", "A21", "A22", "A23", gap=0.15)],
    note=note(("eyebrow", "B2.1 · setup"), ("title", "One column per group"),
              ("text", "B = Napa (dummy = 1)\nC = Other (dummy = 0)")))
V.sheet([
    "In B21, COUNT of column U gives 287 Napa wines. In C21, COUNT of column V gives 519 other wines.",
    "287 plus 519 is 806, so no wine is lost.",
], [T("B21"), T("C21")],
    note=note(("eyebrow", "B2.1 · group sizes"), ("title", "n₁ = 287, n₂ = 519"),
              ("code", "=COUNT(Data!$U$2:$U$807)\n=COUNT(Data!$V$2:$V$807)"),
              ("text", "Check: 287 + 519 = 806.")))
V.sheet([
    "In row 22, AVERAGE of each column. A Napa wine costs $43.06 on average; the other wines, $22.13.",
], [T("B22"), P("C22", select=True)],
    note=note(("eyebrow", "B2.1 · group means"), ("title", "Mean price by group"),
              ("code", "=AVERAGE(Data!$U$2:$U$807)\n=AVERAGE(Data!$V$2:$V$807)"),
              ("kv", [("Napa", "$43.06"), ("Other", "$22.13")])))
V.sheet([
    "In row 23, the variance, with VAR.S. The variance is the standard deviation squared. Napa: 2,690. Others: 171.",
    "The Napa variance is about 16 times bigger. The groups clearly do not have equal variances, and that is why "
    "the project asks for Welch's t-test, which does not assume equal variances.",
], [T("B23"), P("C23", select=True), Sync(1), callout("B23:C23", ORANGE, "≈ 16 times bigger", "right")],
    note=note(("eyebrow", "B2.1 · group variances"), ("title", "Very unequal variances"),
              ("code", "=VAR.S(Data!$U$2:$U$807)\n=VAR.S(Data!$V$2:$V$807)"),
              ("kv", [("Napa s²", "2,690"), ("Other s²", "171"), ("ratio", "≈ 15.7")]),
              ("text", "Lecture 10: when variances differ, pooling is not justified. When in doubt, use Welch.")))
V.sheet([
    "Step one, the hypotheses. H₀: the mean price of Napa wines equals the mean price of other wines. H₁: the two "
    "means are different.",
], [clear_callouts(), P("A25")],
    note=note(("eyebrow", "Step 1 · hypotheses"), ("title", "H₀: μ(Napa) = μ(other)\nH₁: μ(Napa) ≠ μ(other)"),
              ("text", "Equivalently H₀: μ(Napa) − μ(other) = 0. Two-sided, α = 0.05.")))
V.sheet([
    "Step two, the test statistic. In B26, the difference in means: B22 minus C22, which is $20.93.",
], [P("A26", "A27", "A28", "A29", "A30", "A31", "A32", gap=0.1), T("B26")],
    note=note(("eyebrow", "Step 2 · test statistic"), ("title", "Difference in means"), ("code", "=B22-C22"),
              ("kv", [("43.06 − 22.13", "$20.93")])))
V.sheet([
    "In B27, the Welch standard error: the square root of B23 over B21, plus C23 over C21. Each group brings its own "
    "variance divided by its own size. It equals 3.11.",
    "In B28, the Welch t: the difference divided by its standard error. t equals 6.72.",
], [T("B27"), Sync(1), T("B28")],
    note=note(("eyebrow", "Step 2 · test statistic"), ("title", "Welch t = (x̄₁ − x̄₂) / √(s₁²/n₁ + s₂²/n₂)"),
              ("code", "=SQRT(B23/B21+C23/C21)\n=B26/B27"),
              ("kv", [("SE (Welch)", "3.11"), ("t", "6.72")])))
V.sheet([
    "In B29, Welch's degrees of freedom. The formula is long, so pause the video and copy it from the formula bar. "
    "You do not need to memorise it. It gives 306.3.",
], [T("B29", cps=34)],
    note=note(("eyebrow", "Step 2 · degrees of freedom"), ("title", "Welch df ≈ 306.3"),
              ("code", "=(B23/B21+C23/C21)^2/((B23/B21)^2/(B21-1)+(C23/C21)^2/(C21-1))"),
              ("text", "The Satterthwaite approximation. Fewer df than n₁ + n₂ − 2 = 804, because the noisy Napa "
                       "group dominates the SE.")))
V.sheet([
    "Step three, the p-value. In B30, T.DIST.2T of the absolute t, with B29 degrees of freedom. The result is "
    "8.89 E−11.",
    ("E−11 means times ten to the minus eleven: 0.0000000000889. Extremely small.",
     "E minus 11 means times ten to the minus eleven. That is zero point, then ten zeros, then 889. Extremely small."),
], [T("B30")],
    note=note(("eyebrow", "Step 3 · p-value"), ("title", "p ≈ 8.9 × 10⁻¹¹"), ("code", "=T.DIST.2T(ABS(B28),B29)"),
              ("text", "Excel writes very small numbers in E notation:\n8.89E-11 = 0.0000000000889")))
V.sheet([
    "Excel also has a built-in t-test function, T.TEST. In B31, give it the two columns, then 2 for a two-tailed "
    "test, then 3 for the type.",
    "Type 3 is Excel's two-sample test with unequal variances, which is Welch's test. It returns the p-value directly, "
    "and it matches B30.",
    "The tiny difference in the third digit is because T.DIST.2T cuts the degrees of freedom down to the whole "
    "number 306, while T.TEST uses 306.3. Both say the same thing.",
], [T("B31")],
    note=note(("eyebrow", "Excel's own t-test"), ("title", "T.TEST(array1, array2, tails, type)"),
              ("code", "=T.TEST(Data!$U$2:$U$807,Data!$V$2:$V$807,2,3)"),
              ("kv", [("tails = 2", "two-sided"), ("type 1", "paired"), ("type 2", "equal variances (pooled)"),
                      ("type 3", "unequal variances (Welch)")])))
V.sheet([
    "Step four, the decision. The p-value is far below 0.05, so we reject H₀.",
    "In wine terms: on average, Napa wines cost about $21 more per bottle than other wines, and a gap this large is "
    "not a sampling accident.",
    "If you use Excel's Data Analysis ToolPak, the tool called t-Test: Two-Sample Assuming Unequal Variances gives "
    "the same t statistic and p-value. We use formulas because the project wants them visible.",
], [T("B32"), Sync(1), callout("B26", BLUE, "$20.93 more", "right")],
    note=note(("eyebrow", "Step 4 · decision"), ("title", "Reject H₀"),
              ("code", '=IF(B30<0.05,"Reject H0: means differ","Do not reject H0")'),
              ("strong", "Napa wines cost about $20.93 more per bottle on average (raw difference)."),
              ("text", "ToolPak route (Excel): Data → Data Analysis → t-Test: Two-Sample Assuming Unequal Variances.")))

# ------------------------------------------------------------------ regression
V.chapter("B2.2 · Regression on the dummy")
V.slide(slide(s_regression, 1), [
    "B2 part 2: regress Price on the Napa dummy alone. The model is Price equals β₀ plus β₁ times Napa, plus error.",
    "Look at what the model predicts. For a wine that is not from Napa, Napa equals 0, so the prediction is just β₀. "
    "For a Napa wine, Napa equals 1, so the prediction is β₀ plus β₁.",
])
V.slide(slide(s_regression, 2), [
    "Least squares makes each prediction the mean of its own group, just like in B1. So β₀ becomes the mean of the "
    "other wines, and β₁ becomes the gap between the two means.",
])
V.slide(slide(s_regression, 3), [
    "On the chart, the regression line simply joins the two group means. It moves over by one, from 0 to 1, so its "
    "slope is the difference: $20.93.",
])
V.sheet([
    "In Excel we run this regression with the LINEST function. Scroll right to column K. LINEST returns a block of "
    "five rows and two columns, so we first select K42 to L46.",
], [clear_callouts(), Scroll(top=38, sx="D"), P("K41")],
    note=note(("eyebrow", "B2.2 · LINEST"), ("title", "Regression with one formula"),
              ("text", "LINEST(Y range, X range, TRUE, TRUE) returns a 5 × 2 block of results."),
              ("text", "Select the whole block first, then type the formula once.")))
V.sheet([
    "Now type equals LINEST, open a bracket. First the Y range, Price: Data B2 to B807. Then the X range, the Napa "
    "column: Data F2 to F807. Then TRUE, and TRUE again, and close the bracket.",
    "The first TRUE keeps the intercept. The second TRUE asks for the full statistics.",
    "Because the formula fills a whole block, press Control, Shift and Enter together. In Excel 365, plain Enter also works.",
], [ArrayFormula("K42:L46", SRC["K42"][0], [SRC[r][1] for r in cells_in("K42:L46")], cps=26), Sync(2)],
    note=note(("eyebrow", "B2.2 · LINEST"), ("title", "Type once, confirm as an array"),
              ("code", "=LINEST(Data!$B$2:$B$807,Data!$F$2:$F$807,TRUE,TRUE)"),
              ("kv", [("Y", "Price, Data!B2:B807"), ("X", "Napa dummy, Data!F2:F807"), ("1st TRUE", "keep intercept"),
                      ("2nd TRUE", "full statistics")]),
              ("text", "Confirm with Ctrl + Shift + Enter (LibreOffice and older Excel).")))
V.sheet([
    "LINEST always uses the same layout. Top row: the slope, then the intercept. Second row: their standard errors.",
    "Third row: R squared, and the standard error of the regression. Fourth row: the F statistic and the residual "
    "degrees of freedom. Fifth row: the regression and residual sums of squares.",
], [callout("K42", BLUE, "slope = β₁", "left"), callout("L42", ORANGE, "intercept = β₀", "right")],
    note=note(("eyebrow", "B2.2 · reading LINEST"), ("title", "The 5 × 2 layout"),
              ("kv", [("row 1", "slope 20.93  |  intercept 22.13"), ("row 2", "SE 2.403  |  SE 1.434"),
                      ("row 3", "R² 0.086  |  s(e) 32.67"), ("row 4", "F 75.87  |  df 804"),
                      ("row 5", "SSR  |  SSE")]),
              ("text", "R² = 0.086: the Napa dummy alone explains about 9% of the variation in Price.")))
V.sheet([
    "Back in columns A and B, the team gave each of these numbers a name, with short formulas that point to the "
    "LINEST block. That makes the output easy to read.",
], [clear_callouts(), Scroll(top=38, sx=0), P("A41", "A42", "A43", "A44", "A45", "A46", "A47", "A48", "A49", "A50",
                                              gap=0.08),
    P("B42", "B43", "B44", "B45", "B46", "B47", "B48", "B49", "B50", gap=0.1, select=True)],
    note=note(("eyebrow", "B2.2 · tidy output"), ("title", "Name every number"),
              ("code", "=$K$44   R Square\n=$L$45   residual df\n=$K$46   SSR\n=$L$46   SSE"),
              ("text", "Same idea as the Regression Statistics block of the ToolPak output.")))
V.sheet([
    "Then, in rows 52 to 54, a coefficient table like the regression output from class.",
    "The coefficient, its standard error, the t statistic as coefficient over standard error, the p-value with "
    "T.DIST.2T and 804 degrees of freedom, and the 95% confidence interval.",
], [P("A52", "B52", "C52", "D52", "E52", "F52", "G52", "H52", "A53", "A54", gap=0.08), T("B53"), T("B54"),
    P("C53", "C54", "D53", "D54", "E53", "E54", "F53", "F54", "G53", "G54", "H53", "H54", gap=0.14, select=True)],
    note=note(("eyebrow", "B2.2 · coefficient table"), ("title", "Coefficient, SE, t, p, CI"),
              ("code", "=$L$42          intercept\n=$K$42          slope\n=B54/C54        t\n"
                       "=T.DIST.2T(ABS(D54),$B$47)"),
              ("text", "t = coefficient ÷ standard error, df = n − 2 = 804.")))
V.sheet([
    "Now compare. The intercept is 22.13. The slope is 20.93.",
    "Scroll up: 22.13 is exactly C22, the mean price of the other wines, and 20.93 is exactly B26, the difference "
    "between the two means.",
], [callout("B53", ORANGE, "intercept 22.13", "right"), callout("B54", BLUE, "slope 20.93", "right"), Sync(1),
    clear_callouts(), Scroll(top=17), callout("C22", ORANGE, "mean other 22.13", "right"),
    callout("B26", BLUE, "difference 20.93", "right")],
    note=note(("eyebrow", "B2.2 · the key result"), ("title", "Slope = difference of means"),
              ("kv", [("intercept β₀", "22.13 = mean other"), ("slope β₁", "20.93 = 43.06 − 22.13")])))
V.sheet([
    "To prove it in the workbook, rows 56 and 57 subtract them. Both checks give exactly zero.",
    "So the slope equals the difference between the two group means, and the intercept equals the mean of the "
    "non-Napa wines.",
    "With the Data Analysis ToolPak, the Regression tool gives the same coefficients. In LibreOffice it is under "
    "Data, Statistics, Regression.",
], [clear_callouts(), Scroll(top=38), T("A56", cps=40, keys=()), T("B56"), T("A57", cps=40, keys=()), T("B57")],
    note=note(("eyebrow", "B2.2 · proof"), ("title", "Check cells = 0"),
              ("code", "=B54-B26\n=B53-C22"),
              ("text", "Excel route: Data → Data Analysis → Regression, Y = B1:B807, X = F1:F807, tick Labels.\n"
                       "LibreOffice: Data → Statistics → Regression.")))

# ------------------------------------------------------------------ pooled
V.chapter("B2.3 · Regression t = pooled t")
V.slide(slide(s_pooled, 1), [
    "B2 part 3. The slope's t statistic is 8.71, but Welch's t was 6.72. The difference in means is the same, "
    "20.93, so the standard errors must be different.",
])
V.slide(slide(s_pooled, 3), [
    "The regression reproduces exactly the pooled two-sample t-test: the one that assumes both groups have the same "
    "variance. Let's prove it in Excel.",
])
V.sheet([
    "In B35, the pooled variance: each group's variance weighted by its degrees of freedom, n minus 1, then divided "
    "by n₁ plus n₂ minus 2. It is 1,067.",
], [Scroll(top=33), P("A34", "A35", "A36", "A37", "A38", "A39", gap=0.1), T("B35", cps=34)],
    note=note(("eyebrow", "B2.3 · pooled t-test"), ("title", "Pooled variance s_p²"),
              ("code", "=((B21-1)*B23+(C21-1)*C23)/(B21+C21-2)"),
              ("text", "One variance for both groups: a weighted average of 2,690 and 171."),
              ("kv", [("s_p²", "1,067.26")])))
V.sheet([
    "In B36, the pooled standard error: the square root of the pooled variance times, one over n₁ plus one over n₂. "
    "It is 2.403, exactly the standard error of the slope.",
], [T("B36"), callout("B36", BLUE, "2.403", "right"), callout("C54", BLUE, None)],
    note=note(("eyebrow", "B2.3 · pooled t-test"), ("title", "Pooled SE = SE of the slope"),
              ("code", "=SQRT(B35*(1/B21+1/C21))"), ("kv", [("pooled SE", "2.403"), ("LINEST SE of slope", "2.403")])))
V.sheet([
    "In B37, the pooled t, which is 8.71, the same as the slope's t. In B38, its degrees of freedom: 804, the same "
    "as the regression's residual degrees of freedom.",
    "And in B39, T.TEST with type 2, the equal-variance test, gives the same p-value as the regression: 1.7 E−17.",
], [T("B37"), callout("B37", BLUE, "8.71", "right"), callout("D54", BLUE, None), T("B38"),
    callout("B38", BLUE, "804", "right"), callout("B47", BLUE, None), Sync(1), T("B39"),
    callout("B39", BLUE, "1.7E-17", "right"), callout("E54", BLUE, None)],
    note=note(("eyebrow", "B2.3 · pooled t-test"), ("title", "Everything matches"),
              ("code", "=B26/B36\n=B21+C21-2\n=T.TEST(Data!$U$2:$U$807,Data!$V$2:$V$807,2,2)"),
              ("kv", [("t", "8.71 = slope t"), ("df", "804 = residual df"), ("p", "1.7E-17 = slope p")])))
V.sheet([
    "Why? Ordinary least squares assumes that every wine's error has one common variance. It estimates that single "
    "variance from all the residuals together.",
    "Check it yourself: the standard error of the regression, in B44, is 32.67. Squared, that is 1,067: the pooled "
    "variance in B35. Pooling the two groups is exactly what the equal-variance t-test does.",
], [clear_callouts(), Sync(1), callout("B44", ORANGE, "32.67² = 1,067", "right"), callout("B35", ORANGE, "1,067", "right")],
    note=note(("eyebrow", "B2.3 · the reason"), ("title", "OLS assumes one error variance"),
              ("text", "Regression residual variance = SSE / (n − 2) = 858,075.6 / 804 = 1,067.26 = s_p²."),
              ("strong", "So the slope's SE is the pooled SE, and its t is the pooled t.")))
V.sheet([
    "Row 58 checks it: slope t minus pooled t equals zero.",
    "Welch's test lets each group keep its own variance. Here the Napa variance is 16 times larger, so Welch's "
    "standard error is the more honest one, and the pooled one is too small. Luckily both tests reject H₀, so the "
    "conclusion does not change.",
], [clear_callouts(), Scroll(top=38), T("A58", cps=40, keys=()), T("B58"), Sync(1), P("A59")],
    note=note(("eyebrow", "B2.3 · answer"), ("title", "Regression = pooled t-test"),
              ("code", "=D54-B37"),
              ("text", "Welch (SE 3.11, t 6.72) is safer here because the variances are so unequal. "
                       "Both reject H₀ at 5%.")))

# ------------------------------------------------------------------ rank test
V.chapter("B2.4 · A rank-based check")
V.slide(slide(s_rank, 0), [
    "B2 part 4. In A5 our team drew the histogram of Price. It is strongly skewed to the right: most wines cost "
    "under $40, but a few cost over $300.",
    "A t-test compares means, and a few very expensive bottles can pull a mean a long way.",
])
V.slide(slide(s_rank, 1), [
    "The safer rank-based test from class is the Wilcoxon rank-sum test, also called the Mann-Whitney U test.",
])
V.slide(slide(s_rank, 2), [
    "It sorts all 806 prices and replaces each one with its rank: 1 for the cheapest, up to 806 for the most "
    "expensive. Then it asks whether Napa wines tend to have higher ranks than the others.",
    ("A $440 bottle simply gets the top rank, so it cannot dominate the result.",
     "A 440-dollar bottle simply gets the top rank, so it cannot dominate the result."),
])
V.slide(slide(s_rank, 3), [
    "The project says you do not need to run it: just name it, and explain why it is safer.",
    "One more point for the Q&A: with 806 wines, the t-tests are still reasonable, because the central limit theorem "
    "makes sample means close to normal. The rank test is a robustness check.",
])

# ------------------------------------------------------------------ report + QA
V.chapter("Writing it up")
V.slide(slide(s_report_b1), [
    "Here is what to write in the report for B1. Each answer needs the key numbers from Excel, then two to four "
    "sentences of interpretation, and for the test, all four steps.",
    "Pause the video and use this as your template.",
])
V.slide(slide(s_report_b2), [
    "And here is B2. Every number points back to a cell in the Part B sheet, so the lecturer can trace it.",
    "Pause here too, and write it in your own words.",
])

V.chapter("Q&A practice")
for page in (0, 1):
    for k in (1, 2, 3):
        q, a = QA[page * 3 + k - 1]
        V.slide(slide(s_qa, page, k), [(q, QA_SAY[page * 3 + k - 1][0]), (a, QA_SAY[page * 3 + k - 1][1])])
V.slide(slide(s_end), [
    "That's it for B1 and B2. Re-watch any chapter you need, then try to rebuild these cells yourself without the "
    "video.",
    "If you can do that, and explain each step out loud, you are ready for the Q&A. Good luck!",
])

if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--preview", nargs="*", type=float)
    ap.add_argument("--out", default=os.path.join(HERE, "out", "B1_B2_guide.mp4"))
    args = ap.parse_args()
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    V.compile()
    if args.preview is not None:
        info = V.render(args.out, preview_times=args.preview, preview_dir=os.path.join(HERE, "preview"))
    else:
        info = V.render(args.out)
    print({k: v for k, v in info.items() if k != "chapters"})
    for s, name in info["chapters"]:
        print(f"{int(s // 60):02d}:{int(s % 60):02d}  {name}")
